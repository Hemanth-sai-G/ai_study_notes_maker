"""Text extraction adapters used during document ingestion."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import shutil

import fitz
from docx import Document as WordDocument
from PIL import Image
from pptx import Presentation


class ExtractionError(ValueError):
    """Raised when a supported file cannot be read as its advertised type."""


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    page_or_slide_count: int | None
    status: str = "ready"
    message: str | None = None


IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp"}


def extract_text(file_bytes: bytes, suffix: str) -> ExtractionResult:
    """Read supported study files without sending them outside the computer."""
    suffix = suffix.lower()
    try:
        if suffix == ".pdf":
            pdf = fitz.open(stream=file_bytes, filetype="pdf")
            pages = [f"[Page {number}]\n{page.get_text('text').strip()}" for number, page in enumerate(pdf, start=1)]
            return ExtractionResult("\n\n".join(part for part in pages if part.strip()), len(pdf))
        if suffix == ".docx":
            document = WordDocument(BytesIO(file_bytes))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            tables = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
            return ExtractionResult("\n".join(paragraphs + tables), None)
        if suffix == ".pptx":
            presentation = Presentation(BytesIO(file_bytes))
            slides: list[str] = []
            for number, slide in enumerate(presentation.slides, start=1):
                slide_text = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                if slide_text:
                    slides.append(f"[Slide {number}]\n" + "\n".join(slide_text))
            return ExtractionResult("\n\n".join(slides), len(presentation.slides))
        if suffix == ".txt":
            for encoding in ("utf-8-sig", "utf-16", "latin-1"):
                try:
                    return ExtractionResult(file_bytes.decode(encoding).strip(), None)
                except UnicodeDecodeError:
                    continue
            raise ExtractionError("The text file uses an unsupported encoding.")
        if suffix in IMAGE_TYPES:
            Image.open(BytesIO(file_bytes)).verify()
            return _extract_image_text(file_bytes)
    except (fitz.FileDataError, OSError, ValueError, KeyError, TypeError) as error:
        raise ExtractionError(f"StudyMate could not read this {suffix[1:].upper()} file.") from error
    raise ExtractionError(f"Files of type {suffix} are not supported.")


def _extract_image_text(file_bytes: bytes) -> ExtractionResult:
    """OCR images only if both the local Python adapter and Tesseract are available."""
    if shutil.which("tesseract") is None:
        return ExtractionResult(
            text="",
            page_or_slide_count=1,
            status="ocr_unavailable",
            message="Image saved successfully. Install local Tesseract OCR before this image can be read.",
        )
    try:
        import pytesseract

        image = Image.open(BytesIO(file_bytes))
        return ExtractionResult(pytesseract.image_to_string(image).strip(), 1)
    except Exception as error:  # OCR failures should not discard a valid user upload.
        return ExtractionResult("", 1, "ocr_unavailable", f"Image saved, but OCR could not finish: {error}")
