import streamlit as st
import tempfile
import os
import fitz  # PyMuPDF
from docx import Document  # python-docx library

class FileProcessor:
    def __init__(self):
        self.whisper_model = None
        
    def load_whisper_model(self):
        if self.whisper_model is None:
            import whisper
            self.whisper_model = whisper.load_model("base")
        return self.whisper_model

    def _read_text_file(self, uploaded_file):
        """Read plain text files with simple encoding fallbacks."""
        raw_bytes = uploaded_file.read()

        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return raw_bytes.decode(encoding).strip()
            except UnicodeDecodeError:
                continue

        raise ValueError("Unable to decode the uploaded text file.")
        
    def process_file(self, uploaded_file):
        """Process uploaded file and extract text"""
        file_extension = uploaded_file.name.split('.')[-1].lower()
        
        try:
            if file_extension == 'pdf':
                return self.extract_text_from_pdf(uploaded_file)
            elif file_extension == 'docx':
                return self.extract_text_from_docx(uploaded_file)
            elif file_extension in ['txt', 'md']:
                return self._read_text_file(uploaded_file)
            elif file_extension in ['mp3', 'wav', 'm4a', 'flac']:
                return self.transcribe_audio(uploaded_file)
            elif file_extension in ['mp4', 'avi', 'mov', 'mkv']:
                return self.transcribe_video(uploaded_file)
            else:
                st.error(f"Unsupported file type: {file_extension}")
                return None
        except Exception as e:
            st.error(f"Error processing file: {str(e)}")
            return None
            
    def extract_text_from_pdf(self, uploaded_file):
        """Extract text from PDF file"""
        text = ""
        with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = tmp_file.name
            
        try:
            doc = fitz.open(tmp_file_path)
            for page in doc:
                text += page.get_text()
            doc.close()
        finally:
            os.unlink(tmp_file_path)
            
        return text.strip()
    
    def extract_text_from_docx(self, uploaded_file):
        """Extract text from DOCX file"""
        text = ""
        with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = tmp_file.name
            
        try:
            doc = Document(tmp_file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Also extract text from tables if present
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\t"
                    text += "\n"
                    
        finally:
            os.unlink(tmp_file_path)
            
        return text.strip()
        
    def transcribe_audio(self, uploaded_file):
        """Transcribe audio file to text"""
        with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{uploaded_file.name.split(".")[-1]}') as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = tmp_file.name
            
        try:
            model = self.load_whisper_model()
            result = model.transcribe(tmp_file_path)
            return result["text"]
        except ImportError:
            st.error("Audio transcription dependencies are not installed. Install openai-whisper to enable audio support.")
            return None
        finally:
            os.unlink(tmp_file_path)
            
    def transcribe_video(self, uploaded_file):
        """Extract audio from video and transcribe"""
        with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{uploaded_file.name.split(".")[-1]}') as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = tmp_file.name
            
        audio_path = tmp_file_path.replace(tmp_file_path.split('.')[-1], 'wav')
        
        try:
            # Extract audio from video
            from moviepy.editor import VideoFileClip

            video = VideoFileClip(tmp_file_path)
            video.audio.write_audiofile(audio_path, verbose=False, logger=None)
            video.close()
            
            # Transcribe audio
            model = self.load_whisper_model()
            result = model.transcribe(audio_path)
            return result["text"]
        except ImportError:
            st.error("Video transcription dependencies are not installed. Install moviepy and openai-whisper to enable video support.")
            return None
        finally:
            if os.path.exists(tmp_file_path):
                os.unlink(tmp_file_path)
            if os.path.exists(audio_path):
                os.unlink(audio_path)
